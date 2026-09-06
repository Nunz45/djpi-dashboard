import json, re, io, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

d = json.load(open('sk-upi-mentah.json', encoding='utf-8'))

# --- sisa nomor baris di depan nama, dibersihkan berulang ---
for b in d:
    s = b['namaJurnal']
    for _ in range(3):
        s2 = re.sub(r'^\s*\d{1,4}[.\s]+', '', s)
        if s2 == s: break
        s = s2
    b['namaJurnal'] = s.strip(' .;:,-')
    b['perluDicek'] = 'ya' if len(b['namaJurnal']) < 5 else ''

def masaBerlaku(b):
    if b['asalAkhir'] == 'tertulis di SK':
        return 'Vol %s No %s Tahun %s s.d. Vol %s No %s Tahun %s' % (
            b['volMulai'], b['noMulai'], b['tahunMulai'], b['volAkhir'], b['noAkhir'], b['tahunAkhir'])
    if b['asalAkhir'].startswith('DITURUNKAN'):
        return 'mulai Vol %s No %s Tahun %s — berakhir sekitar %s (diturunkan, tidak tertulis di SK)' % (
            b['volMulai'], b['noMulai'], b['tahunMulai'], b['tahunAkhir'])
    return 'tidak tercantum di SK'

for b in d:
    b['masaBerlaku'] = masaBerlaku(b)

# --- catatan yang berlaku = SK terbaru per e-ISSN; sisanya jadi riwayat ---
terbaru = {}
for b in sorted(d, key=lambda x: x['urut']):
    terbaru[b['eIssn']] = b
berlaku = sorted(terbaru.values(), key=lambda x: (-x['urut'], x['namaJurnal'].lower()))
for b in berlaku:
    b['jumlahSk'] = sum(1 for x in d if x['eIssn'] == b['eIssn'])

MARUN = 'FF7F0000'; KERTAS = 'FFF4EFEB'; EMAS = 'FFFBF3DF'
tepi = Border(*[Side(style='thin', color='FFD8D0C8')] * 4)

def tulisSheet(ws, kolom, data, lebar):
    for i, k in enumerate(kolom, 1):
        c = ws.cell(1, i, k)
        c.font = Font(bold=True, color='FFFFFFFF', size=10)
        c.fill = PatternFill('solid', fgColor=MARUN)
        c.alignment = Alignment(vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = lebar[i - 1]
    ws.row_dimensions[1].height = 30
    for r, row in enumerate(data, 2):
        for i, v in enumerate(row, 1):
            c = ws.cell(r, i, v)
            c.alignment = Alignment(vertical='top', wrap_text=(i in (1, 3)))
            c.border = tepi
            c.font = Font(size=10)
        if r % 2 == 0:
            for i in range(1, len(kolom) + 1):
                ws.cell(r, i).fill = PatternFill('solid', fgColor=KERTAS)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = 'A1:%s%d' % (get_column_letter(len(kolom)), len(data) + 1)

wb = Workbook()

# ---------- Sheet 1: yang berlaku ----------
ws = wb.active; ws.title = 'Masa Berlaku'
kolom = ['Nama Jurnal', 'Nomor SK', 'Masa Berlaku', 'e-ISSN', 'Peringkat',
         'Jenis Penetapan', 'Periode SK', 'Tanggal SK', 'Sumber Tanggal Akhir',
         'Jumlah SK', 'Perlu Dicek', 'Berkas Sumber']
data = [[b['namaJurnal'], b['nomorSk'], b['masaBerlaku'], b['eIssn'],
         ('SINTA ' + b['peringkat']) if b['peringkat'] else '',
         b['jenis'], b['periode'], b['tanggalSk'], b['asalAkhir'],
         b['jumlahSk'], b['perluDicek'], b['berkas']] for b in berlaku]
tulisSheet(ws, kolom, data, [44, 22, 46, 12, 11, 22, 18, 17, 26, 10, 11, 40])
# tandai baris yang akhirnya diturunkan / tidak ada
for r, b in enumerate(berlaku, 2):
    if b['asalAkhir'] != 'tertulis di SK':
        for i in (3, 9):
            ws.cell(r, i).fill = PatternFill('solid', fgColor=EMAS)

# ---------- Sheet 2: riwayat lengkap ----------
ws2 = wb.create_sheet('Riwayat Semua SK')
kolom2 = ['Nama Jurnal', 'Nomor SK', 'Masa Berlaku', 'e-ISSN', 'Peringkat',
          'Jenis Penetapan', 'Periode SK', 'Tanggal SK', 'Sumber Tanggal Akhir', 'Status']
d2 = sorted(d, key=lambda x: (x['eIssn'], x['urut']))
data2 = []
for b in d2:
    st = 'BERLAKU' if terbaru.get(b['eIssn']) is b else 'digantikan SK berikutnya'
    data2.append([b['namaJurnal'], b['nomorSk'], b['masaBerlaku'], b['eIssn'],
                  ('SINTA ' + b['peringkat']) if b['peringkat'] else '',
                  b['jenis'], b['periode'], b['tanggalSk'], b['asalAkhir'], st])
tulisSheet(ws2, kolom2, data2, [44, 22, 46, 12, 11, 22, 18, 17, 26, 24])

# ---------- Sheet 3: catatan ----------
ws3 = wb.create_sheet('Catatan & Keterbatasan')
ws3.column_dimensions['A'].width = 120
catatan = [
 ('JUDUL', 'Daftar Jurnal UPI Menurut SK Akreditasi/Reakreditasi SINTA'),
 ('', ''),
 ('H', 'Cara berkas ini dibuat'),
 ('', '1. SK diunduh dari tautan pada "Daftar SK Akreditasi dan Reakreditasi SINTA.md", disimpan di'),
 ('', '   C:\\Users\\Asus\\Documents\\SK Akreditasi bersama SK yang sudah Anda punya sebelumnya.'),
 ('', '2. Teks lampiran diurai dengan PyMuPDF. Jangkarnya e-ISSN, bukan nama jurnal, karena hanya'),
 ('', '   e-ISSN yang bentuknya seragam di seluruh periode.'),
 ('', '3. Baris disaring dengan mencocokkan penerbit ke "Universitas Pendidikan Indonesia".'),
 ('', '   Pencocokan dilakukan setelah spasi dinormalkan, karena PDF memecah isi sel jadi'),
 ('', '   beberapa baris sehingga pencarian biasa menghasilkan negatif palsu.'),
 ('', ''),
 ('H', 'Tiga format lampiran, dan akibatnya pada kolom Masa Berlaku'),
 ('', 'SK 2018      : hanya memuat peringkat dan penerbit. TIDAK ADA rentang volume sama sekali.'),
 ('', 'SK 2019-2020 : memuat "mulai Volume X Nomor Y Tahun Z" saja, tanpa "sampai".'),
 ('', '               Tahun berakhirnya DITURUNKAN (awal + 5 tahun) dan ditandai kuning.'),
 ('', 'SK 2021 ke atas: memuat rentang penuh "mulai ... sampai ...". Ini yang paling tepercaya.'),
 ('', ''),
 ('H', 'Yang perlu diperhatikan sebelum dipakai mengambil keputusan'),
 ('', '- Baris berlatar kuning berarti tanggal akhirnya TIDAK tertulis di SK. Jangan dipakai'),
 ('', '  menghitung tenggat sebelum diperiksa ke sertifikat jurnal yang bersangkutan.'),
 ('', '- SK memuat volume dan nomor, bukan bulan. Bulan pastinya baru bisa dihitung kalau'),
 ('', '  frekuensi terbit jurnal diketahui.'),
 ('', '- Kolom "Perlu Dicek" bertanda ya berarti penguraian nama gagal dan harus dibaca manual.'),
 ('', '- Sheet "Riwayat Semua SK" memperlihatkan seluruh penetapan per jurnal, termasuk yang'),
 ('', '  sudah digantikan, supaya perpindahan peringkat terlihat.'),
 ('', ''),
 ('H', 'SK yang belum berhasil didapat'),
 ('', '- 2018 Periode III (34/E/KPT/2018): tautan acuan mengarah ke Scribd, bukan PDF.'),
 ('', '- 2022 Periode I (105/E/KPT/2022): cermin repositori menolak unduhan.'),
 ('', '- 2025 Periode II (295/C/C3/KPT/2026 dan 156/C/C3/KPT/2026): acuan hanya memberi tautan'),
 ('', '  pengumuman, bukan PDF. Anda punya berkasnya di folder SK Akreditasi tetapi teksnya'),
 ('', '  belum bisa diurai; perlu diperiksa manual.'),
 ('', '- 2021 Periode II dan 2019 Periode III terunduh, tetapi tidak menghasilkan baris UPI.'),
 ('', '  Perlu dicek apakah memang tidak ada jurnal UPI di sana, atau lampirannya berupa gambar.'),
]
r = 1
for jenis, t in catatan:
    c = ws3.cell(r, 1, t)
    if jenis == 'JUDUL':
        c.font = Font(bold=True, size=14, color=MARUN)
    elif jenis == 'H':
        c.font = Font(bold=True, size=11, color=MARUN)
    else:
        c.font = Font(size=10)
    r += 1

TUJUAN = r'C:\Users\Asus\Documents\Daftar Jurnal UPI - Masa Berlaku Akreditasi SINTA.xlsx'
wb.save(TUJUAN)
print('Excel ditulis: %s' % TUJUAN)
print('  Sheet "Masa Berlaku"      : %d jurnal' % len(berlaku))
print('  Sheet "Riwayat Semua SK"  : %d penetapan' % len(d2))
print()
print('  akhir tertulis di SK      : %d' % sum(1 for b in berlaku if b['asalAkhir'] == 'tertulis di SK'))
print('  akhir diturunkan          : %d' % sum(1 for b in berlaku if b['asalAkhir'].startswith('DITURUNKAN')))
print('  akhir tidak ada           : %d' % sum(1 for b in berlaku if b['asalAkhir'] == 'tidak ada'))
print('  nama perlu dicek manual   : %d' % sum(1 for b in berlaku if b['perluDicek']))
