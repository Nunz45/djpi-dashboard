# Cek silang: hasil penguraian SK  <->  Sheet1 direktori DJPI,
# diperkaya dengan lembar tambahan berisi hasil reakreditasi periode 2 tahun 2025.
#
# Kunci gabung utama adalah e-ISSN yang dinormalkan (tanpa tanda hubung, X huruf
# besar). Nama jurnal hanya dipakai sebagai jaring kedua, dan setiap kecocokan
# yang HANYA lewat nama ditandai supaya bisa diperiksa manusia.
import csv, io, json, re, os, subprocess, datetime
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ID_UTAMA   = '180PvnhYAvplEkJXoeQJBHMkwjUlAX4nWytucslHOYLc'
ID_TAMBAHAN= '1hC1yrhFyjn8QXaXJwEmppl5dGg7yNqNuZY_Q12h3hkc'
TMP = os.path.dirname(os.path.abspath(__file__))

def unduhCsv(sid, nama, gid=None):
    url = 'https://docs.google.com/spreadsheets/d/%s/export?format=csv' % sid
    if gid is not None: url += '&gid=%s' % gid
    p = os.path.join(TMP, nama)
    subprocess.run(['curl','-sS','-L','--max-time','120','-o',p,url], check=True)
    with io.open(p, encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))

def issn(v):
    v = re.sub(r'[^0-9Xx]', '', str(v or '')).upper()
    return v if re.match(r'^\d{7}[\dX]$', v) else ''

# Kolom ISSN di Sheet1: AF = E-ISSN, AG = P-ISSN. Kolom lama "ISSN" (X) tidak
# dipakai lagi. e-ISSN jadi kunci utama karena lampiran SK mencatat EISSN;
# p-ISSN hanya jaring kedua dan kecocokannya SELALU ditandai, sebab mencocokkan
# nomor cetak ke kolom elektronik bisa menghasilkan pasangan yang keliru.
def eissnBaris(r):
    return issn(r.get('E-ISSN'))

def pissnBaris(r):
    return issn(r.get('P-ISSN'))

def kunci(n):
    n = re.sub(r'\([^)]*\)', ' ', n or '')
    n = re.sub(r'[^a-z0-9 ]', ' ', n.lower())
    return re.sub(r'\s+', ' ', n).strip()

# Cerminan issuePerTahun_ di Code.js:5603 — dipakai menurunkan bulan dari nomor terbitan.
def issuePerTahun(s):
    s = (s or '').upper()
    if not s.strip(): return None
    if re.search(r'BULANAN|MONTHLY', s): return 12
    if re.search(r'DWI\s*BULAN|BIMONTH', s): return 6
    if re.search(r'TRIWULAN|KUARTAL|QUARTER', s): return 4
    if re.search(r'TENGAH\s*TAHUN|SEMESTER|SEMIANNUAL|DUA\s*KALI|2\s*KALI', s): return 2
    if re.search(r'TAHUNAN|ANNUAL|SEKALI\s*SETAHUN|SETAHUN\s*SEKALI', s): return 1
    n = re.sub(r'[^0-9]', '', s)
    n = int(n) if n else 0
    return n if 1 <= n <= 24 else None

RENTANG = re.compile(r'mulai\s+Volume\s+(\d+)\s+Nomor\s+(\d+)\s+Tahun\s+(\d{4})\s+'
                     r'sampai\s+Volume\s+(\d+)\s+Nomor\s+(\d+)\s+Tahun\s+(\d{4})', re.I)
PERINGKAT_KE = re.compile(r'ke\s+Peringkat\s+(\d)', re.I)
PERINGKAT_DI = re.compile(r'(?:di|Baru)\s+Peringkat\s+(\d)', re.I)

print('Mengunduh Sheet1 dan lembar tambahan...')
utama = unduhCsv(ID_UTAMA, 'utama.csv', '0')
tambahan = unduhCsv(ID_TAMBAHAN, 'tambahan.csv')
sk = json.load(open(os.path.join(TMP, 'sk-upi-mentah.json'), encoding='utf-8'))
print('  Sheet1 %d baris | tambahan %d baris | hasil urai SK %d penetapan'
      % (len(utama), len(tambahan), len(sk)))

# ---------- perkaya: lembar tambahan jadi penetapan SK ----------
# Semua barisnya reakreditasi dengan rentang mulai 2024-2025, cocok dengan
# periode 2 tahun 2025. Nomor SK-nya DIDUGA 295/C/C3/KPT/2026 dan ditandai
# supaya tidak dipakai sebagai bukti resmi sebelum PDF-nya diperiksa.
# Idempoten: buang dulu baris pengayaan dari jalannya yang lalu, kalau tidak
# menjalankan ulang skrip ini akan menggandakan sembilan penetapan itu.
sk = [x for x in sk if x.get('berkas') != 'lembar tambahan Google Sheets']
baru = 0
for r in tambahan:
    e = issn(r.get('ISSN'))
    ket = r.get('Peringkat') or ''
    m = RENTANG.search(ket)
    pk = PERINGKAT_KE.search(ket) or PERINGKAT_DI.search(ket)
    sk.append({
        'urut': 2025.2, 'periode': '2025 Periode II',
        'nomorSk': '295/C/C3/KPT/2026 [PERLU KONFIRMASI]',
        'tanggalSk': '2 Januari 2026', 'berkas': 'lembar tambahan Google Sheets',
        'namaJurnal': (r.get('Nama Jurnal') or '').strip(), 'eIssn': e,
        'penerbitCuplikan': (r.get('Nama Penerbit') or '')[:120],
        'jenis': 'Reakreditasi', 'peringkat': pk.group(1) if pk else '',
        'volMulai': m.group(1) if m else '', 'noMulai': m.group(2) if m else '',
        'tahunMulai': m.group(3) if m else '',
        'volAkhir': m.group(4) if m else '', 'noAkhir': m.group(5) if m else '',
        'tahunAkhir': m.group(6) if m else '',
        'asalAkhir': 'tertulis di SK' if m else 'tidak ada',
    })
    baru += 1
print('  ditambahkan dari lembar tambahan: %d penetapan' % baru)

# ---------- rekam yang berlaku per jurnal ----------
terbaru = {}
for b in sorted(sk, key=lambda x: x['urut']):
    if b['eIssn']: terbaru[b['eIssn']] = b
petaNama = {}
for e, b in terbaru.items(): petaNama.setdefault(kunci(b['namaJurnal']), b)

# ---------- cek silang ----------
TAHUN_INI = datetime.date.today().year

# e-ISSN kembar di Sheet1: dua jurnal berbeda memakai nomor yang sama. Ini galat
# data yang harus diketahui, karena e-ISSN adalah kunci gabung seluruh berkas ini.
hitungIssn = {}
for r in utama:
    e = eissnBaris(r)
    if e: hitungIssn.setdefault(e, []).append((r.get('NAMA JURNAL') or '').strip())
issnKembar = {e: v for e, v in hitungIssn.items() if len(v) > 1}
hasil, ringkas = [], {'cocok-issn':0,'cocok-pissn':0,'cocok-nama':0,'tanpa-sk':0}
for r in utama:
    nama = (r.get('NAMA JURNAL') or '').strip()
    if not nama: continue
    e = eissnBaris(r)
    pi = pissnBaris(r)
    status = (r.get('STATUS AKREDITASI') or '').strip()
    terakreditasi = status.upper().startswith('SINTA')
    pkSheet = re.sub(r'[^0-9]', '', status) if terakreditasi else ''

    rec, caraGabung = (terbaru.get(e), 'e-ISSN (AF)') if e and e in terbaru else (None, '')
    if not rec and pi and pi in terbaru:
        rec, caraGabung = terbaru[pi], 'p-ISSN (AG)'
    if not rec:
        rec = petaNama.get(kunci(nama))
        caraGabung = 'nama jurnal' if rec else ''
    if rec:
        ringkas['cocok-issn' if caraGabung.startswith('e-ISSN') else
                ('cocok-pissn' if caraGabung.startswith('p-ISSN') else 'cocok-nama')] += 1
    else:
        ringkas['tanpa-sk'] += 1

    catatan = []
    if not rec and terakreditasi:
        catatan.append('Sheet1 menyebut %s tetapi tidak ada di SK mana pun' % status)
    if rec and not terakreditasi:
        catatan.append('SK menyebut terakreditasi tetapi Sheet1 menulis "%s"' % (status or 'kosong'))
    if rec and pkSheet and rec['peringkat'] and pkSheet != rec['peringkat']:
        # Kalau rekam SK yang ketemu sudah tua, penyebab paling mungkin bukan
        # Sheet1 yang salah melainkan SK terbarunya belum masuk kumpulan ini.
        if rec['urut'] < 2023:
            catatan.append('peringkat beda: Sheet1 %s vs SK %s — tetapi rekam SK dari %s, '
                           'kemungkinan SK terbarunya belum terkumpul'
                           % (pkSheet, rec['peringkat'], rec['periode']))
        else:
            catatan.append('peringkat beda: Sheet1 %s vs SK %s (rekam SK sudah baru, '
                           'perlu diadu ke sertifikat)' % (pkSheet, rec['peringkat']))
    if not e:
        catatan.append('e-ISSN (AF) kosong di Sheet1')
    if caraGabung == 'p-ISSN (AG)':
        catatan.append('cocok lewat p-ISSN, bukan e-ISSN — lampiran SK mencatat e-ISSN, '
                       'jadi pasangan ini perlu diperiksa')
    if e and e in issnKembar:
        catatan.append('e-ISSN KEMBAR di Sheet1, dipakai juga oleh: %s'
                       % ', '.join(n for n in issnKembar[e] if n != nama)[:60])

    # bandingkan tahun kedaluwarsa
    thSheet = ''
    for sumber in (r.get('TANGGAL EXPIRED'), r.get('MASA BERLAKU SK AKREDITASI')):
        m = re.search(r'\b(20\d{2})\b', str(sumber or ''))
        if m: thSheet = m.group(1); break
    thSk = rec['tahunAkhir'] if rec else ''
    if thSheet and thSk and thSheet != thSk:
        catatan.append('tahun berakhir beda: Sheet1 %s vs SK %s' % (thSheet, thSk))

    # pemeriksaan bernilai tertinggi: nomor akhir vs frekuensi terbit
    bulan = ''
    ipt = issuePerTahun(r.get('ISSUE'))
    if rec and rec.get('noAkhir') and ipt:
        no = int(rec['noAkhir'])
        if no > ipt:
            catatan.append('nomor akhir %d melebihi %d terbitan/tahun — penurunan bulan tidak sah' % (no, ipt))
        else:
            bulan = '%s-%02d' % (thSk, max(1, min(12, round(12 * no / ipt))))

    hasil.append({
        'nama': nama, 'kluster': (r.get('KLUSTER') or '').strip(), 'eIssn': e, 'pIssn': pi,
        'statusSheet1': status, 'peringkatSk': ('SINTA ' + rec['peringkat']) if (rec and rec['peringkat']) else '',
        'nomorSk': rec['nomorSk'] if rec else '', 'periodeSk': rec['periode'] if rec else '',
        'masaBerlaku': ('Vol %s No %s Thn %s s.d. Vol %s No %s Thn %s' % (
            rec['volMulai'], rec['noMulai'], rec['tahunMulai'],
            rec['volAkhir'], rec['noAkhir'], rec['tahunAkhir'])) if (rec and rec['asalAkhir'] == 'tertulis di SK') else (
            rec['asalAkhir'] if rec else ''),
        'tahunAkhirSk': thSk, 'tanggalExpiredSheet1': (r.get('TANGGAL EXPIRED') or '').strip(),
        'issue': (r.get('ISSUE') or '').strip(), 'terbitPerTahun': ipt or '',
        'perkiraanBulanBerakhir': bulan,
        'caraGabung': caraGabung or 'tidak ketemu',
        'catatan': ' | '.join(catatan),
    })

# jurnal di SK yang tidak ada di Sheet1
adaDiSheet1 = {x for r in utama for x in (eissnBaris(r), pissnBaris(r)) if x}
namaSheet1 = {kunci((r.get('NAMA JURNAL') or '')) for r in utama}
yatim = [b for e, b in terbaru.items()
         if e not in adaDiSheet1 and kunci(b['namaJurnal']) not in namaSheet1]

print()
print('=== HASIL CEK SILANG ===')
print('  cocok lewat e-ISSN (kolom AF)  : %d' % ringkas['cocok-issn'])
print('  cocok lewat p-ISSN (kolom AG)  : %d  (ditandai, perlu diperiksa)' % ringkas['cocok-pissn'])
print('  cocok hanya lewat nama jurnal : %d  (perlu diperiksa manusia)' % ringkas['cocok-nama'])
print('  tidak ketemu di SK            : %d' % ringkas['tanpa-sk'])
print('  ada di SK tetapi tidak di Sheet1: %d' % len(yatim))
print()
for tag in ['tidak ada di SK mana pun','Sheet1 menulis','peringkat beda','kemungkinan SK terbarunya belum terkumpul','perlu diadu ke sertifikat','tahun berakhir beda','KEMBAR','melebihi']:
    n = sum(1 for h in hasil if tag in h['catatan'])
    print('  %-42s %d' % (tag, n))
print('  perkiraan bulan berakhir berhasil dihitung : %d' % sum(1 for h in hasil if h['perkiraanBulanBerakhir']))

json.dump({'hasil': hasil, 'yatim': yatim}, open(os.path.join(TMP,'cek-silang.json'),'w',encoding='utf-8'),
          ensure_ascii=False, indent=2)
json.dump(sk, open(os.path.join(TMP,'sk-upi-mentah.json'),'w',encoding='utf-8'), ensure_ascii=False, indent=2)
print()
print('Ditulis: cek-silang.json dan sk-upi-mentah.json (sudah termasuk lembar tambahan)')
